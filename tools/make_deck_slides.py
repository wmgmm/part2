#!/usr/bin/env python3
"""Add the Matt Mort exercise slides to Matt Hayden's deck, in his style.

Reads his deck, never writes back to it. Slide 21 ("Gemini notebook", his
placeholder for Exercise 05) is the template: its sidebar, logos and title
geometry are cloned for the new slides, and its rough-notes text box is
replaced in place. Output is a new file.

    python3 tools/make_deck_slides.py "in.pptx" "out.pptx"
"""
import copy
import io
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt

TEMPLATE = 21          # 1-based: "Gemini notebook"
COPILOT_LOGO = (12, "Content Placeholder 4")   # 960x960 Copilot mark
CANVAS_ICON = (18, "Picture 5")                # 960x960 Canvas mark
DIM = RGBColor(0x59, 0x59, 0x5B)

# One line each, drawn from the site's step titles. [OPTIONAL] marks the
# stretch items, as Matt asked. "Website: Exercise 04" maps slide to page.
SLIDES = [
    dict(title="Skills", caption="Copilot + Gemini", icon=COPILOT_LOGO, website=None, bullets=[
        "A skill is a workflow written down: a small markdown file.",
        "Attach it in Copilot or Gemini, or paste it into Notebook.",
        "Repeatable, shareable, and your judgement is in the file.",
        "Five in today's pack, on the exercise pages.",
        "[OPTIONAL] Write your own, from any task you repeat.",
    ]),
    dict(title="The Module", caption="Gemini or Copilot", icon=COPILOT_LOGO, website="Exercise 04", bullets=[
        "Download the training skill and the plan from the website.",
        "Paste the prompt, attach both files, send.",
        "Review and augment: the gap note is where to look.",
        "[OPTIONAL] Run it in the other tool and compare.",
        "Responsible AI: the human value is in the review.",
    ]),
    dict(title=None, caption=None, icon="keep", website="Exercise 05", bullets=[
        "New notebook. Add your training PDF, or the site's backup.",
        "Slide deck: paste Cardiff_Brand_SKILL.md as the description.",
        "Generate. Notebook builds the deck to the skill.",
        "[OPTIONAL] Second run with TheMattsBrandSkill.md. Compare.",
        "Responsible AI: internal only. Public-facing is a human job.",
    ]),
    dict(title="The Dashboard", caption="Gemini Canvas", icon=CANVAS_ICON, website="Exercise 06", bullets=[
        "New chat, Canvas on, attach the CSV, paste the build prompt.",
        "Open ours and run yours beside it. Do they agree?",
        "Verify and repair in the same chat, reattaching the CSV.",
        "[OPTIONAL] Cleaner audit: a new chat, your HTML and the CSV.",
        "Responsible AI: check the numbers before you tell the story.",
    ]),
]


def shape_by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def clone_slide(prs, src):
    """New slide on the same layout with every shape of src copied, pictures re-linked."""
    new = prs.slides.add_slide(src.slide_layout)
    for ph in list(new.shapes):
        ph._element.getparent().remove(ph._element)
    for sh in src.shapes:
        el = copy.deepcopy(sh._element)
        new.shapes._spTree.append(el)
    # Re-point every image relationship at a relationship on the new slide part.
    for blip in new.shapes._spTree.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
        rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        if rid:
            target = src.part.rels[rid].target_part
            blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed",
                     new.part.relate_to(target, RT.IMAGE))
    return new


def move_slide(prs, slide, index):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    el = ids[prs.slides.index(slide)]
    lst.remove(el)
    lst.insert(index, el)


def set_title(slide, text):
    t = shape_by_name(slide, "Title 1")
    para = t.text_frame.paragraphs[0]
    run = para.runs[0] if para.runs else para.add_run()
    run.text = text
    for extra in para.runs[1:]:
        extra._r.getparent().remove(extra._r)


def set_caption(slide, text):
    box = shape_by_name(slide, "TextBox 20")
    para = box.text_frame.paragraphs[0]
    run = para.runs[0]
    run.text = text
    run.font.size = Pt(36 if len(text) <= 9 else 22)


def swap_icon(prs, slide, source):
    """Replace the sidebar icon (Picture 8) with a picture from another slide."""
    old = shape_by_name(slide, "Picture 8")
    left, top, width, height = old.left, old.top, old.width, old.height
    old._element.getparent().remove(old._element)
    if source is None:
        return
    src_slide = prs.slides[source[0] - 1]
    blob = shape_by_name(src_slide, source[1]).image.blob
    slide.shapes.add_picture(io.BytesIO(blob), left, top, width, height)


def set_bullets(slide, bullets, website):
    box = shape_by_name(slide, "TextBox 1")
    box.top, box.height = Inches(1.9), Inches(4.9)
    tf = box.text_frame
    tf.word_wrap = True
    for para in list(tf.paragraphs)[1:]:
        para._p.getparent().remove(para._p)
    first = tf.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    first.text = ""
    for i, text in enumerate(bullets):
        para = first if i == 0 else tf.add_paragraph()
        optional = text.startswith("[OPTIONAL]")
        run = para.add_run()
        run.text = text
        run.font.size = Pt(20)
        run.font.name = "Aptos"
        if optional:
            run.font.color.rgb = DIM
        para.space_after = Pt(8)
    if website:
        para = tf.add_paragraph()
        run = para.add_run()
        run.text = f"Website: {website}   https://wmgmm.github.io/part2/"
        run.font.size = Pt(14)
        run.font.name = "Aptos"
        run.font.color.rgb = DIM
        para.space_before = Pt(10)


def main(src_path, out_path):
    prs = Presentation(src_path)
    template = prs.slides[TEMPLATE - 1]
    made = []
    for spec in SLIDES:
        if spec["title"] is None:
            slide = template                       # Exercise 05: edit slide 21 in place
        else:
            slide = clone_slide(prs, template)
            set_title(slide, spec["title"])
            set_caption(slide, spec["caption"])
            if spec["icon"] != "keep":
                swap_icon(prs, slide, spec["icon"])
        set_bullets(slide, spec["bullets"], spec["website"])
        made.append(slide)
    # Order: Skills, The Module, Gemini notebook (the original 21), The Dashboard.
    move_slide(prs, made[0], TEMPLATE - 1)
    move_slide(prs, made[1], TEMPLATE)
    move_slide(prs, made[3], TEMPLATE + 2)
    prs.save(out_path)
    print(f"wrote {out_path}: {len(prs.slides)} slides")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
