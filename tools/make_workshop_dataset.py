#!/usr/bin/env python3
"""Build the Exercise 06 workshop dataset from the raw HESA extract.

Run from the repo root:  python3 tools/make_workshop_dataset.py

WHY THIS EXISTS
    Exercise 06 failed in a real run on 2026-09-05. Gemini replied: "I cannot run
    code on this file. Additionally, I cannot see the contents of
    HESA_Estates_Management.xlsx." Researched the same day, the causes are:

    1. 346,050 rows cannot fit any Gemini context window. Google publishes
       32k tokens with no AI plan, 128k on AI Plus, 1M on Pro/Ultra
       (support.google.com/gemini/answer/16275805). A long-format table that
       size is tens of millions of tokens. Reading it was never possible, which
       is exactly what "I cannot see the contents" means.
    2. Google's code execution tool documents CSV and text, NOT xlsx. The Vertex
       AI file-input MIME list is .cpp .csv .java .jpeg .js .png .py .ts .xml,
       and the API docs say "Code execution works best with text and CSV files"
       (ai.google.dev/gemini-api/docs/code-execution). Google's own Gemini
       Notebook accepts csv and refuses xlsx. So the upload succeeded and the
       only tool that could have done the work would not take the file.
    3. The 13 MB was never the problem: Google's documented limit is 100 MB.

    Hence: a small, wide, CSV-first dataset. The .xlsx is written too, because
    Microsoft Copilot DOES list .xlsx for data analysis, and because the second
    sheet of questions is useful to a human opening it in Excel.

SHAPE
    The source is long format: one row per provider/year/metric, five stacked
    HESA tables, ~200 categories. Every model has to pivot that to wide before
    it can answer anything. This script does the pivot once, so the workshop
    can be about the analysis rather than about schema archaeology.
"""

import csv
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "public" / "placeholders" / "HESA_Estates_Management.xlsx"
OUT = ROOT / "public" / "placeholders"
STEM = "HESA_Estates_Workshop"

# HESA category -> our column. Note Energy uses HESA's own published total:
# summing the components would DOUBLE COUNT, because the component list holds
# both "Natural gas used as input for a CHP unit" and "Electricity consumed
# from onsite CHP".
FIELDS = {
    "Total number of buildings": "Buildings",
    "Total gross internal area (m2)": "FloorArea_m2",
    "Total energy consumption (kWh)": "Energy_kWh",
    "Total water consumption (m3)": "Water_m3",
    "Total number of car parking spaces": "CarSpaces",
    "Total number of cycle spaces": "CycleSpaces",
    "Total percentage of renewable energy purchased through green tariffs (%)": "RenewablesPct",
    # Not in the original brief. Added deliberately: Exercise 06 is about
    # emissions, and the facilitator answer keys are all carbon figures.
    "Total scope 1 and 2 carbon emissions (Kg CO2e)": "Scope12_tCO2e",
}

COLUMNS = [
    "Institution", "AcademicYear", "Buildings", "FloorArea_m2", "Energy_kWh",
    "Water_m3", "CarSpaces", "CycleSpaces", "RenewablesPct", "Scope12_tCO2e",
]

# CarSpaces, CycleSpaces and RenewablesPct appear in two HESA tables. Pin each
# field to one table so a provider can never get two different answers.
TABLE_FOR = {
    "Buildings": "Table-1", "FloorArea_m2": "Table-1",
    "CarSpaces": "Table-1", "CycleSpaces": "Table-1",
    "Energy_kWh": "Table-2", "Water_m3": "Table-2", "RenewablesPct": "Table-2",
    "Scope12_tCO2e": "Table-3",
}

RUSSELL_GROUP = {
    "The University of Birmingham", "The University of Bristol",
    "The University of Cambridge", "Cardiff University", "University of Durham",
    "The University of Edinburgh", "The University of Exeter",
    "The University of Glasgow",
    "Imperial College of Science, Technology and Medicine",
    "King's College London", "The University of Leeds",
    "The University of Liverpool",
    "London School of Economics and Political Science",
    "The University of Manchester", "Newcastle University",
    "University of Nottingham", "The University of Oxford",
    "Queen Mary University of London", "Queen's University Belfast",
    "The University of Sheffield", "The University of Southampton",
    "University College London", "The University of Warwick",
    "The University of York",
}

# Every Welsh provider in the record. Cardiff University is in both sets and is
# counted once. Matched on exact name: "Cardiff" alone would also catch Cardiff
# Metropolitan, which is a different institution with a different estate.
WELSH = {
    "Aberystwyth University", "Bangor University",
    "Cardiff Metropolitan University", "Cardiff University",
    "Swansea University", "University of South Wales",
    "University of Wales Trinity Saint David", "Wrexham University",
}

KEEP = RUSSELL_GROUP | WELSH

# HESA's legal names are not what anyone calls these places. The brief asks for
# friendly names, so drop a leading "The" and give Imperial the name it uses.
RENAME = {"Imperial College of Science, Technology and Medicine": "Imperial College London"}

QUESTIONS = [
    "1. Which universities appear most sustainable?",
    "2. How has Cardiff University's energy consumption changed over time?",
    "3. Which institutions have the highest renewable energy usage?",
    "4. Which universities appear to be outliers?",
    "5. Create a one-page executive summary for a Vice-Chancellor.",
    "6. Identify three universities Cardiff should benchmark against.",
    "7. What trends are visible across the sector?",
    "8. Which institutions have improved most over the last five years?",
]


def to_number(raw, field):
    """HESA stores every value as text, and percentages carry a % sign."""
    if raw is None:
        return None
    s = str(raw).strip().replace(",", "").rstrip("%").strip()
    if not s or not re.fullmatch(r"-?\d+(\.\d+)?", s):
        return None
    v = float(s)
    if field == "Scope12_tCO2e":      # HESA publishes kg; the keys are tonnes
        v /= 1000.0
    return round(v, 1) if v % 1 else int(v)


def main():
    if not SRC.exists():
        raise SystemExit(f"Source missing: {SRC}")

    ws = openpyxl.load_workbook(SRC, read_only=True)["Sheet1"]
    rows = ws.iter_rows(values_only=True)
    for _ in range(10):                       # HESA's metadata block
        next(rows)
    head = next(rows)
    ix = {k: n for n, k in enumerate(head)}

    table = {}
    for r in rows:
        provider = r[ix["HE Provider"]]
        if provider not in KEEP:
            continue
        field = FIELDS.get(r[ix["Category"]])
        if not field or r[ix["Table"]] != TABLE_FOR[field]:
            continue
        key = (RENAME.get(provider, re.sub(r"^The ", "", provider)),
               r[ix["Academic Year"]])
        table.setdefault(key, {})[field] = to_number(r[ix["Value"]], field)

    records = [
        [inst, year] + [table[(inst, year)].get(c) for c in COLUMNS[2:]]
        for inst, year in sorted(table)
    ]

    with (OUT / f"{STEM}.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for rec in records:
            w.writerow(["" if v is None else v for v in rec])

    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Institution_Year_Summary"
    sheet.append(COLUMNS)
    for rec in records:
        sheet.append(rec)                     # None writes a real blank cell
    q = wb.create_sheet("Workshop_Questions")
    q.append(["Questions to try"])
    for line in QUESTIONS:
        q.append([line])
    q.append([])
    # Said out loud because question 7 invites a claim this file cannot support.
    q.append(["Note: this is 31 UK universities, the Russell Group plus every "
              "Welsh provider, not the whole sector."])
    q.append(["Adapted from HESA, www.hesa.ac.uk, CC BY 4.0. Filtered and "
              "reshaped from the published Estates Management record."])
    wb.save(OUT / f"{STEM}.xlsx")

    insts = sorted({r[0] for r in records})
    years = sorted({r[1] for r in records})
    print(f"  rows          {len(records)}")
    print(f"  institutions  {len(insts)}")
    print(f"  years         {len(years)}  ({years[0]} to {years[-1]})")
    for n, col in enumerate(COLUMNS[2:], start=2):
        filled = sum(1 for r in records if r[n] is not None)
        print(f"    {col:16} {filled:4}/{len(records)}  {100*filled/len(records):5.1f}%")
    for name, path in ((f"{STEM}.csv", OUT / f"{STEM}.csv"),
                       (f"{STEM}.xlsx", OUT / f"{STEM}.xlsx")):
        print(f"  {name:32} {path.stat().st_size/1024:8.1f} KB")


if __name__ == "__main__":
    main()
